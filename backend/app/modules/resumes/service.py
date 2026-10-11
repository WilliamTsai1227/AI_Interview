"""履歷：建立（帶 upload_id）、列表、主要履歷、刪除；worker 解析（≤ 10 頁、60 秒、並行上限）。"""
import asyncio
import hashlib
import io
from uuid import UUID

import structlog

from app.ai import client as ai
from app.ai.agents import parsing
from app.core import db, redis, storage
from app.core.config import settings
from app.core.errors import AppError, not_found
from app.modules.uploads import service as uploads
from app.workers.queue import enqueue, tx

log = structlog.get_logger()
_parse_sem = asyncio.Semaphore(4)


def out(r, detail: bool = False) -> dict:
    d = {"id": str(r["id"]), "filename": r["filename"], "size_bytes": r["size_bytes"], "is_primary": r["is_primary"],
         "parse_status": r["parse_status"], "created_at": r["created_at"]}
    if detail:
        parsed = r["parsed_json"] or None
        d.update({"parse_error": r["parse_error"], "parsed": parsed, "analysis": r["analysis_json"],
                  "parsed_at": r["updated_at"] if r["parse_status"] == "parsed" else None})
    return d


async def create(user_id: UUID, upload_id: UUID, set_primary: bool | None) -> tuple[int, dict]:
    await redis.rate_limit("resume_upload", str(user_id))
    filename = await uploads.filename_of(upload_id) or "履歷.pdf"
    async with tx() as conn:
        f = await uploads.consume(conn, user_id, upload_id, "resume")
        data = await storage.read(f["object_key"])
        digest = hashlib.sha256(data).hexdigest()
        dup = await conn.fetchrow(
            """
            SELECT * FROM resumes WHERE user_id = $1 AND sha256 = $2 AND deleted_at IS NULL
               AND created_at > now() - interval '7 days'
            """, user_id, digest)
        if dup:
            return 200, out(dup)
        has_primary = await conn.fetchval(
            "SELECT 1 FROM resumes WHERE user_id = $1 AND is_primary AND deleted_at IS NULL", user_id)
        primary = set_primary if set_primary is not None else not has_primary
        if primary:
            await conn.execute("UPDATE resumes SET is_primary = false WHERE user_id = $1 AND is_primary", user_id)
        r = await conn.fetchrow(
            """
            INSERT INTO resumes (user_id, filename, mime_type, size_bytes, file_key, sha256, is_primary)
            VALUES ($1, $2, $3, $4, $5, $6, $7) RETURNING *
            """, user_id, filename, f["content_type"], f["size"], f["object_key"], digest, primary)
        await enqueue(conn, "parse_resume", {"resume_id": str(r["id"])}, dedupe_key=f"parse_resume:{r['id']}")
    return 202, out(r)


async def list_(user_id: UUID) -> list[dict]:
    async with db.connection() as conn:
        rs = await conn.fetch(
            "SELECT * FROM resumes WHERE user_id = $1 AND deleted_at IS NULL ORDER BY is_primary DESC, created_at DESC",
            user_id)
    return [out(r) for r in rs]


async def get(user_id: UUID, resume_id: UUID) -> dict:
    async with db.connection() as conn:
        r = await conn.fetchrow("SELECT * FROM resumes WHERE id = $1 AND user_id = $2 AND deleted_at IS NULL",
                                resume_id, user_id)
    if not r:
        raise not_found("找不到這份履歷")
    return out(r, detail=True)


async def patch(user_id: UUID, resume_id: UUID, is_primary: bool | None, filename: str | None) -> dict:
    async with tx() as conn:
        r = await conn.fetchrow(
            "SELECT * FROM resumes WHERE id = $1 AND user_id = $2 AND deleted_at IS NULL FOR UPDATE", resume_id, user_id)
        if not r:
            raise not_found("找不到這份履歷")
        if is_primary:
            await conn.execute("UPDATE resumes SET is_primary = false WHERE user_id = $1 AND is_primary AND id <> $2",
                               user_id, resume_id)
            await conn.execute("UPDATE resumes SET is_primary = true WHERE id = $1", resume_id)
        if filename:
            await conn.execute("UPDATE resumes SET filename = $2 WHERE id = $1", resume_id, filename.strip()[:200])
        r = await conn.fetchrow("SELECT * FROM resumes WHERE id = $1", resume_id)
    return out(r, detail=True)


async def delete(user_id: UUID, resume_id: UUID) -> None:
    async with tx() as conn:
        r = await conn.fetchrow(
            "UPDATE resumes SET deleted_at = now(), is_primary = false WHERE id = $1 AND user_id = $2 AND deleted_at IS NULL RETURNING *",
            resume_id, user_id)
        if not r:
            raise not_found("找不到這份履歷")
        if r["is_primary"]:
            await conn.execute(
                """
                UPDATE resumes SET is_primary = true WHERE id = (
                  SELECT id FROM resumes WHERE user_id = $1 AND deleted_at IS NULL ORDER BY created_at DESC LIMIT 1)
                """, user_id)
        await conn.execute("UPDATE target_jobs SET resume_id = NULL WHERE resume_id = $1", resume_id)
    try:
        await storage.delete(r["file_key"])
    except Exception as e:
        log.warning("resume_file_delete_failed", error=str(e))


# ---------- 給其他模組使用 ----------
async def primary_id(conn: db.Connection, user_id: UUID) -> UUID | None:
    return await conn.fetchval("SELECT id FROM resumes WHERE user_id = $1 AND is_primary AND deleted_at IS NULL", user_id)


async def owned(conn: db.Connection, user_id: UUID, resume_id: UUID) -> db.Record:
    r = await conn.fetchrow("SELECT * FROM resumes WHERE id = $1 AND user_id = $2 AND deleted_at IS NULL", resume_id, user_id)
    if not r:
        raise not_found("找不到這份履歷")
    return r


def snapshot(r: db.Record | None) -> dict | None:
    """出題與評分需要的履歷欄位（resume_snapshot）。"""
    if not r or r["parse_status"] != "parsed" or not r["parsed_json"]:
        return None
    p = r["parsed_json"]
    return {"resume_id": str(r["id"]), "sha256": r["sha256"], "summary": p.get("summary"),
            "experiences": p.get("experiences", [])[:6], "skills": p.get("skills", [])[:20],
            "achievements": p.get("achievements", [])[:8], "projects": p.get("projects", [])[:5]}


# ---------- worker ----------
def _extract_text(data: bytes, ext: str) -> str:
    if ext == "pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in reader.pages[: settings.resume_max_pages])
    if ext == "docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)
    raise ValueError("舊版 Word（.doc）無法解析，請另存成 PDF 或 .docx 再上傳")


async def handle_parse(payload: dict) -> None:
    rid = UUID(payload["resume_id"])
    async with db.connection() as conn:
        r = await conn.fetchrow("UPDATE resumes SET parse_status = 'parsing' WHERE id = $1 AND deleted_at IS NULL RETURNING *", rid)
    if not r:
        return
    try:
        async with _parse_sem:
            data = await storage.read(r["file_key"])
            ext = {"application/pdf": "pdf", "application/msword": "doc"}.get(r["mime_type"], "docx")
            raw = await asyncio.wait_for(asyncio.to_thread(_extract_text, data, ext), timeout=60)
            if len(raw.strip()) < 30:
                raise ValueError("讀不到履歷文字（可能是掃描檔），請上傳可以選取文字的 PDF")
            parsed = await parsing.parse_resume(raw, user_id=r["user_id"])
            emb = await ai.embed(raw[:6000], user_id=r["user_id"])
    except Exception as e:
        async with db.connection() as conn:
            await conn.execute("UPDATE resumes SET parse_status = 'failed', parse_error = $2 WHERE id = $1",
                               rid, str(e)[:300])
        log.warning("resume_parse_failed", resume_id=str(rid), error=str(e))
        return
    p = parsed.model_dump()
    analysis = {"highlights": p.pop("highlights"), "improvements": p.pop("improvements")}
    async with tx() as conn:
        await conn.execute(
            """
            UPDATE resumes SET parse_status = 'parsed', parsed_text = $2, parsed_json = $3, analysis_json = $4,
                   embedding = $5, parse_error = NULL WHERE id = $1
            """, rid, raw[:50000], p, analysis, emb)
        # 第一份履歷：補到還沒有預設履歷的目標職缺，並重新產生面試建議（題組不覆蓋）
        targets = await conn.fetch(
            """
            UPDATE target_jobs SET resume_id = $2
             WHERE user_id = $1 AND resume_id IS NULL AND status = 'active'
               AND EXISTS (SELECT 1 FROM resumes WHERE id = $2 AND is_primary)
             RETURNING id
            """, r["user_id"], rid)
        for t in targets:
            prep = await conn.fetchrow("SELECT id FROM interview_preps WHERE target_job_id = $1", t["id"])
            if prep:
                from app.modules.preps import service as preps
                await preps.start_generation(conn, prep["id"], "analyze", {"resume_id": str(rid)}, quiet=True)
